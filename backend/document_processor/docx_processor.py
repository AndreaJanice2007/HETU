"""DOCX paragraph, table, and embedded-image extraction."""

from __future__ import annotations

from io import BytesIO

from document_processor.errors import CorruptedDocxError
from document_processor.schemas import PreparedImage, ProcessedDocument

MAX_IMAGES = 8


def process_docx(data: bytes, filename: str) -> ProcessedDocument:
    try:
        from docx import Document
        from docx.opc.exceptions import PackageNotFoundError
    except ImportError as exc:
        raise CorruptedDocxError("python-docx is not installed") from exc

    try:
        document = Document(BytesIO(data))
    except PackageNotFoundError as exc:
        raise CorruptedDocxError("DOCX could not be opened (file may be corrupted)") from exc
    except Exception as exc:
        raise CorruptedDocxError("DOCX could not be opened (file may be corrupted)") from exc

    try:
        return _extract(document, filename)
    except CorruptedDocxError:
        raise
    except Exception as exc:
        raise CorruptedDocxError("DOCX could not be processed (file may be corrupted)") from exc


def _extract(document, filename: str) -> ProcessedDocument:
    blocks: list[str] = []
    notes: list[str] = []

    for paragraph in document.paragraphs:
        text = (paragraph.text or "").strip()
        if not text:
            continue
        style = (paragraph.style.name if paragraph.style else "") or ""
        if style.lower().startswith("heading"):
            blocks.append(f"[{style}] {text}")
        else:
            blocks.append(text)

    for table_index, table in enumerate(document.tables, start=1):
        blocks.append(f"--- table {table_index} ---")
        for row in table.rows:
            cells = [" ".join(cell.text.split()) for cell in row.cells]
            blocks.append(" | ".join(cells))

    images = _embedded_images(document)
    if images:
        notes.append(f"Extracted {len(images)} embedded image(s)")
    if not blocks and not images:
        notes.append("DOCX contained no paragraphs, tables, or images")

    return ProcessedDocument(
        kind="docx",
        filename=filename,
        text="\n".join(blocks).strip(),
        images=images[:MAX_IMAGES],
        page_count=1,
        notes=notes,
    )


def _embedded_images(document) -> list[PreparedImage]:
    images: list[PreparedImage] = []
    rels = getattr(getattr(document, "part", None), "rels", {}) or {}
    for index, rel in enumerate(rels.values(), start=1):
        reltype = (getattr(rel, "reltype", "") or "").lower()
        if "image" not in reltype:
            continue
        part = getattr(rel, "target_part", None)
        blob = getattr(part, "blob", None)
        if not blob:
            continue
        content_type = (getattr(part, "content_type", "") or "image/png").lower()
        mime = content_type if content_type.startswith("image/") else "image/png"
        ext = "png"
        if "jpeg" in mime or "jpg" in mime:
            mime = "image/jpeg"
            ext = "jpg"
        elif "png" in mime:
            mime = "image/png"
            ext = "png"
        images.append(
            PreparedImage(
                filename=f"docx-image-{index}.{ext}",
                mime=mime,
                data=bytes(blob),
                reason="embedded_image",
            )
        )
        if len(images) >= MAX_IMAGES:
            break
    return images
