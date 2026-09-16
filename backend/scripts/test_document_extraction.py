"""Offline tests for HETU document extraction. Uses fictional Adult A only."""

from __future__ import annotations

import io
import sys
from pathlib import Path
from unittest.mock import patch

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from fastapi.testclient import TestClient
from PIL import Image, ImageDraw

from app.main import app
from document_processor.errors import (
    CorruptedDocxError,
    CorruptedPdfError,
    UnreadableImageError,
    UnsupportedFileTypeError,
)
from document_processor.extractor import prepare_document
from document_processor.schemas import ClinicalDocument

FICTIONAL_NOTE = (
    "FICTIONAL DEMO — not a real patient.\n"
    "Adult A, age 54, female.\n"
    "Diagnosis: type 2 diabetes mellitus.\n"
    "Medication: metformin 500 mg twice daily.\n"
    "Lab: HbA1c 7.2% (fictional).\n"
    "Finding: no chest pain today.\n"
)


def _pdf_bytes(text: str) -> bytes:
    import pymupdf

    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((72, 72), text, fontsize=11)
    data = doc.tobytes()
    doc.close()
    return data


def _scanned_pdf_bytes() -> bytes:
    import pymupdf

    image = Image.new("RGB", (400, 200), "white")
    draw = ImageDraw.Draw(image)
    draw.text((20, 80), "Adult A scanned fictional BP 142/88", fill="black")
    buf = io.BytesIO()
    image.save(buf, format="PNG")
    png = buf.getvalue()

    doc = pymupdf.open()
    page = doc.new_page(width=400, height=200)
    page.insert_image(page.rect, stream=png)
    data = doc.tobytes()
    doc.close()
    return data


def _docx_bytes(text: str) -> bytes:
    from docx import Document

    document = Document()
    document.add_heading("Fictional clinic note", level=1)
    document.add_paragraph(text)
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Test"
    table.cell(0, 1).text = "Value"
    table.cell(1, 0).text = "HbA1c"
    table.cell(1, 1).text = "7.2%"
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


def _png_bytes(text: str) -> bytes:
    image = Image.new("RGB", (480, 160), "white")
    draw = ImageDraw.Draw(image)
    draw.text((16, 70), text, fill="black")
    buf = io.BytesIO()
    image.save(buf, format="PNG")
    return buf.getvalue()


def _jpeg_bytes(text: str) -> bytes:
    image = Image.new("RGB", (480, 160), "white")
    draw = ImageDraw.Draw(image)
    draw.text((16, 70), text, fill="black")
    buf = io.BytesIO()
    image.save(buf, format="JPEG", quality=85)
    return buf.getvalue()


def test_processors() -> None:
    pdf = prepare_document(_pdf_bytes(FICTIONAL_NOTE), "adult-a-note.pdf")
    assert pdf.kind == "pdf"
    assert "metformin" in pdf.text.lower()
    assert "type 2 diabetes" in pdf.text.lower()

    scanned = prepare_document(_scanned_pdf_bytes(), "adult-a-scan.pdf")
    assert scanned.kind == "pdf"
    assert scanned.images, "scanned PDF should send page images to the vision model"
    assert scanned.images[0].reason == "scanned_page"

    docx = prepare_document(_docx_bytes(FICTIONAL_NOTE), "adult-a-note.docx")
    assert docx.kind == "docx"
    assert "metformin" in docx.text.lower()
    assert "HbA1c" in docx.text

    png = prepare_document(_png_bytes("Adult A fictional glucose 118"), "glucose.png")
    assert png.kind == "png"
    assert len(png.images) == 1
    assert png.images[0].reason == "original_image"

    jpeg = prepare_document(_jpeg_bytes("Adult A fictional potassium 4.1"), "potassium.jpg")
    assert jpeg.kind == "jpeg"
    print("processors: PASS")


def test_errors() -> None:
    try:
        prepare_document(b"not a document", "note.txt")
        raise AssertionError("expected unsupported type")
    except UnsupportedFileTypeError:
        pass
    try:
        prepare_document(b"%PDF-broken", "broken.pdf")
        raise AssertionError("expected corrupted PDF")
    except CorruptedPdfError:
        pass
    try:
        prepare_document(b"PK\x03\x04not-a-docx", "broken.docx")
        raise AssertionError("expected corrupted DOCX")
    except CorruptedDocxError:
        pass
    try:
        prepare_document(b"\x89PNG\r\n\x1a\nnot-an-image", "broken.png")
        raise AssertionError("expected unreadable image")
    except UnreadableImageError:
        pass
    print("error handling: PASS")


def _auth_header() -> dict:
    client = TestClient(app)
    login = client.post(
        "/api/login",
        json={"email": "arjun.patel@hetu.demo", "password": "ArjunFortis44"},
    )
    assert login.status_code == 200, login.text
    return {"X-User-Id": str(login.json()["user"]["id"])}


def test_extract_endpoint_mocked() -> None:
    client = TestClient(app)
    headers = _auth_header()
    clinical = ClinicalDocument(
        diagnoses=["type 2 diabetes mellitus"],
        medications=["metformin 500 mg twice daily"],
        patient={"name": "Adult A", "age": 54, "sex": "female"},
    )

    def fake_openai(processed):
        return clinical

    with patch("document_processor.extractor.call_openai_extractor", side_effect=fake_openai):
        response = client.post(
            "/api/documents/extract",
            headers=headers,
            files={"file": ("adult-a-note.pdf", _pdf_bytes(FICTIONAL_NOTE), "application/pdf")},
        )
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["clinical"]["diagnoses"] == ["type 2 diabetes mellitus"]
    assert body["processing"]["detected_type"] == "pdf"
    assert body["processing"]["contradiction_engine"] == "not_invoked"
    print("extract endpoint (mocked OpenAI): PASS")


def test_extract_error_statuses() -> None:
    client = TestClient(app)
    headers = _auth_header()
    unsupported = client.post(
        "/api/documents/extract",
        headers=headers,
        files={"file": ("notes.txt", b"hello", "text/plain")},
    )
    assert unsupported.status_code == 415
    assert unsupported.json()["detail"]["error"] == "unsupported_file_type"

    bad_pdf = client.post(
        "/api/documents/extract",
        headers=headers,
        files={"file": ("broken.pdf", b"%PDF-broken", "application/pdf")},
    )
    assert bad_pdf.status_code == 400
    assert bad_pdf.json()["detail"]["error"] == "corrupted_pdf"
    print("extract HTTP errors: PASS")


def test_live_openai_extract() -> None:
    from document_processor.extractor import extract_clinical_document

    data = _pdf_bytes(FICTIONAL_NOTE)
    clinical, meta = extract_clinical_document(data, "adult-a-fictional.pdf")
    print("live model:", meta["model"])
    print("live diagnoses:", clinical.diagnoses)
    print("live medications:", clinical.medications)
    print("live uncertainties:", clinical.uncertainties)
    assert isinstance(clinical.diagnoses, list)
    print("live OpenAI extract: PASS")


def main() -> int:
    test_processors()
    test_errors()
    test_extract_endpoint_mocked()
    test_extract_error_statuses()
    print("document extraction pipeline: PASS")
    if "--live" in sys.argv:
        test_live_openai_extract()
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
